import re
from functools import lru_cache
from io import BytesIO

import fitz
from docx import Document
from sentence_transformers import SentenceTransformer
from sklearn.metrics.pairwise import cosine_similarity


MODEL_NAME = "sentence-transformers/all-MiniLM-L6-v2"

SEMANTIC_WEIGHT = 0.60
SKILL_WEIGHT = 0.25
EXPERIENCE_WEIGHT = 0.10
EDUCATION_WEIGHT = 0.05

SKILLS_LIST = [
    "python", "java", "c", "c++", "javascript", "typescript",
    "html", "css", "react", "angular", "vue", "node.js", "express",
    "mongodb", "sql", "mysql", "postgresql", "sqlite",
    "machine learning", "deep learning", "artificial intelligence",
    "nlp", "natural language processing", "data science", "data analysis",
    "data preprocessing", "data visualization",
    "pandas", "numpy", "scikit-learn", "tensorflow", "pytorch", "keras",
    "flask", "django", "fastapi", "rest api",
    "git", "github", "docker", "kubernetes",
    "aws", "azure", "gcp", "cloud",
    "excel", "power bi", "tableau",
    "opencv", "computer vision",
    "llm", "generative ai", "langchain",
    "communication", "problem solving", "teamwork", "leadership",
]

SKILL_ALIASES = {
    "ml": "machine learning",
    "machine learning": "machine learning",
    "ai": "artificial intelligence",
    "artificial intelligence": "artificial intelligence",
    "nlp": "natural language processing",
    "natural language processing": "natural language processing",
    "js": "javascript",
    "javascript": "javascript",
    "nodejs": "node.js",
    "node.js": "node.js",
    "reactjs": "react",
    "react.js": "react",
    "postgres": "postgresql",
    "ms excel": "excel",
}

EDUCATION_KEYWORDS = [
    "b.tech",
    "b.e",
    "m.tech",
    "bca",
    "mca",
    "computer science",
    "degree",
    "certification",
    "certified",
    "diploma",
    "engineering",
]


def extract_text_from_pdf(file_bytes):
    """Extract text from a PDF resume using PyMuPDF."""
    text = ""

    with fitz.open(stream=file_bytes, filetype="pdf") as pdf_document:
        for page in pdf_document:
            text += page.get_text()

    return text


def extract_text_from_docx(file_bytes):
    """Extract text from DOCX paragraphs and tables."""
    document = Document(BytesIO(file_bytes))
    text_parts = []

    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            text_parts.append(paragraph.text.strip())

    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    text_parts.append(cell.text.strip())

    return " ".join(text_parts)


def extract_resume_text(file_name, file_bytes):
    """Choose the correct text extraction method based on file extension."""
    lower_file_name = file_name.lower()

    if lower_file_name.endswith(".pdf"):
        return extract_text_from_pdf(file_bytes)

    if lower_file_name.endswith(".docx"):
        return extract_text_from_docx(file_bytes)

    return ""


def clean_text(text):
    """Clean text so matching and embeddings receive consistent input."""
    text = text.lower()
    text = re.sub(r"\s+", " ", text)
    text = re.sub(r"[^a-z0-9\s.,+#-]", "", text)
    return text.strip()


def normalize_keyword_text(text):
    """Normalize keywords for simple exact keyword matching."""
    text = text.lower()
    text = text.replace("-", " ")
    text = re.sub(r"\s+", " ", text)
    return text.strip()


def keyword_found(text, keyword):
    """Check whether a full skill or keyword exists in the text."""
    searchable_text = normalize_keyword_text(text)
    searchable_keyword = normalize_keyword_text(keyword)
    pattern = r"(?<![a-z0-9+#])" + re.escape(searchable_keyword) + r"(?![a-z0-9+#])"
    return re.search(pattern, searchable_text) is not None


@lru_cache(maxsize=1)
def load_model():
    """Load the Sentence Transformer model once and reuse it for all API calls."""
    try:
        # Offline-first keeps the final demo working after the model is cached.
        return SentenceTransformer(MODEL_NAME, local_files_only=True)
    except Exception:
        # First-time setup can still download the model when internet is available.
        return SentenceTransformer(MODEL_NAME)


def split_text_into_chunks(text, chunk_size=1200, overlap=200):
    """Split long text into overlapping chunks."""
    words = text.split()
    chunks = []

    start = 0
    while start < len(words):
        end = start + chunk_size
        chunks.append(" ".join(words[start:end]))
        start += chunk_size - overlap

    return chunks


def calculate_semantic_score(job_description, resume_text):
    """Calculate the best semantic similarity between the JD and resume chunks."""
    model = load_model()
    job_embedding = model.encode([job_description])
    resume_chunks = split_text_into_chunks(resume_text)

    if not resume_chunks:
        return 0

    chunk_embeddings = model.encode(resume_chunks)
    similarities = cosine_similarity(job_embedding, chunk_embeddings)[0]
    score = float(max(similarities) * 100)
    score = max(0, min(100, score))
    return round(score, 2)


def extract_skills_from_text(text):
    """Find skills and aliases, then return sorted canonical skill names."""
    found_skills = set()

    for skill in SKILLS_LIST:
        if keyword_found(text, skill):
            found_skills.add(SKILL_ALIASES.get(skill, skill))

    for alias, canonical_skill in SKILL_ALIASES.items():
        if keyword_found(text, alias):
            found_skills.add(canonical_skill)

    return sorted(found_skills)


def calculate_skill_match_score(required_skills, resume_skills):
    """Calculate skill score as matched required skills divided by total skills."""
    if not required_skills:
        return 0

    matched_skills = set(required_skills).intersection(set(resume_skills))
    score = (len(matched_skills) / len(required_skills)) * 100
    return round(score, 2)


def get_missing_skills(required_skills, resume_skills):
    """Return required skills from the JD that are missing in the resume."""
    return sorted(set(required_skills).difference(set(resume_skills)))


def extract_experience_years(text, use_range_upper_value=True):
    """Extract experience from patterns like 1 year, 3+ years, 1-2 years, 1.5 years, yrs."""
    text = text.lower()

    if "fresher" in text or "entry level" in text:
        return 0

    experience_values = []

    pattern = r"(\d+(?:\.\d+)?)\s*-\s*(\d+(?:\.\d+)?)\s*(?:years?|yrs?)|\b(\d+(?:\.\d+)?)\s*\+\s*(?:years?|yrs?)|\b(\d+(?:\.\d+)?)\s*(?:years?|yrs?)"

    for match in re.finditer(pattern, text):
        range_start = match.group(1)
        range_end = match.group(2)
        plus_year = match.group(3)
        single_year = match.group(4)

        if range_start and range_end:
            if use_range_upper_value:
                experience_values.append(float(range_end))
            else:
                experience_values.append(float(range_start))
        elif plus_year:
            experience_values.append(float(plus_year))
        elif single_year:
            experience_values.append(float(single_year))

    if not experience_values:
        return None

    return max(experience_values)


def extract_candidate_name(text, file_name):
    """Try to extract candidate name from first few resume lines."""
    lines = [line.strip() for line in text.splitlines() if line.strip()]

    for line in lines[:6]:
        clean_line = re.sub(r"[^A-Za-z\s]", "", line).strip()
        words = clean_line.split()

        if 2 <= len(words) <= 4 and all(word[:1].isupper() for word in words):
            return clean_line

    return file_name.rsplit(".", 1)[0].replace("_", " ").replace("-", " ").title()


def calculate_experience_score(required_experience, candidate_experience):
    """Score candidate experience compared with required experience."""
    if required_experience is None or candidate_experience is None:
        return 50

    if candidate_experience >= required_experience:
        return 100

    score = (candidate_experience / required_experience) * 100
    return round(score, 2)


def extract_education_keywords(text):
    """Find education and certification keywords in text."""
    found_keywords = []

    for keyword in EDUCATION_KEYWORDS:
        if keyword_found(text, keyword):
            found_keywords.append(keyword)

    return found_keywords


def calculate_education_score(required_keywords, resume_keywords):
    """Score overlap between JD and resume education/certification keywords."""
    if not required_keywords:
        return 70

    matched_keywords = set(required_keywords).intersection(set(resume_keywords))
    score = (len(matched_keywords) / len(required_keywords)) * 100
    return round(score, 2)


def calculate_final_score(
    semantic_score,
    skill_score,
    experience_score,
    education_score,
):
    """Combine all scores using the required hybrid formula."""
    final_score = (
        SEMANTIC_WEIGHT * semantic_score
        + SKILL_WEIGHT * skill_score
        + EXPERIENCE_WEIGHT * experience_score
        + EDUCATION_WEIGHT * education_score
    )

    return round(final_score, 2)


def get_recommendation(final_score):
    """Return the recommendation label based on the final score."""
    if final_score >= 75:
        return "Shortlist"

    if final_score >= 50:
        return "Maybe Review"

    return "Not Recommended"


def get_score_interpretation(final_score):
    """Explain the final score range in simple words."""
    if final_score >= 80:
        return "Excellent Match"

    if final_score >= 60:
        return "Good Match"

    if final_score >= 40:
        return "Average Match"

    return "Low Match"


def get_recommendation_reason(final_score, recommendation):
    """Create a short reason for the recommendation."""
    interpretation = get_score_interpretation(final_score)

    if recommendation == "Shortlist":
        return (
            f"This resume is recommended because the final score is {final_score:.2f}%, "
            f"which is an {interpretation}."
        )

    if recommendation == "Maybe Review":
        return (
            f"This resume needs manual review because the final score is {final_score:.2f}%, "
            f"which is a {interpretation}."
        )

    return (
        f"This resume is not recommended because the final score is {final_score:.2f}%, "
        f"which is a {interpretation}."
    )


def format_list_for_response(items):
    """Convert a list into readable text for JSON and the frontend table."""
    if not items:
        return "None"

    return ", ".join(items)


def score_single_resume(
    file_name,
    file_bytes,
    cleaned_job_description,
    required_skills,
    required_experience,
    required_education_keywords,
):
    """Calculate all ranking scores for one resume."""
    resume_text = extract_resume_text(file_name, file_bytes)
    candidate_name = extract_candidate_name(resume_text, file_name)
    cleaned_resume_text = clean_text(resume_text)

    if not cleaned_resume_text:
        raise ValueError("Could not extract readable text from this resume.")

    semantic_score = calculate_semantic_score(cleaned_job_description, cleaned_resume_text)

    resume_skills = extract_skills_from_text(cleaned_resume_text)
    matched_skills = sorted(set(required_skills).intersection(set(resume_skills)))
    missing_skills = get_missing_skills(required_skills, resume_skills)
    skill_score = calculate_skill_match_score(required_skills, resume_skills)

    candidate_experience = extract_experience_years(
        cleaned_resume_text,
        use_range_upper_value=True,
    )
    experience_score = calculate_experience_score(required_experience, candidate_experience)

    resume_education_keywords = extract_education_keywords(cleaned_resume_text)
    education_score = calculate_education_score(
        required_education_keywords,
        resume_education_keywords,
    )

    final_score = calculate_final_score(
        semantic_score,
        skill_score,
        experience_score,
        education_score,
    )
    recommendation = get_recommendation(final_score)

    return {
        "candidate_name": candidate_name,
        "file_name": file_name,
        "final_score": final_score,
        "semantic_score": semantic_score,
        "skill_score": skill_score,
        "experience_score": experience_score,
        "education_score": education_score,
        "matched_skills": format_list_for_response(matched_skills),
        "missing_skills": format_list_for_response(missing_skills),
        "recommendation": recommendation,
        "reason": get_recommendation_reason(final_score, recommendation),
    }


def rank_resume_files(job_description, resume_files):
    """Rank uploaded resumes and return API-ready results plus a summary."""
    cleaned_job_description = clean_text(job_description)
    required_skills = extract_skills_from_text(cleaned_job_description)
    required_experience = extract_experience_years(
        cleaned_job_description,
        use_range_upper_value=False,
    )
    required_education_keywords = extract_education_keywords(cleaned_job_description)
    results = []

    for resume_file in resume_files:
        try:
            result = score_single_resume(
                resume_file["file_name"],
                resume_file["file_bytes"],
                cleaned_job_description,
                required_skills,
                required_experience,
                required_education_keywords,
            )
            results.append(result)
        except Exception as error:
            file_name = resume_file["file_name"]
            results.append(
                {
                    "candidate_name": file_name.rsplit(".", 1)[0]
                    .replace("_", " ")
                    .replace("-", " ")
                    .title(),
                    "file_name": file_name,
                    "final_score": 0,
                    "semantic_score": 0,
                    "skill_score": 0,
                    "experience_score": 0,
                    "education_score": 0,
                    "matched_skills": "None",
                    "missing_skills": format_list_for_response(required_skills),
                    "recommendation": "Not Recommended",
                    "reason": str(error),
                }
            )

    results = sorted(results, key=lambda item: item["final_score"], reverse=True)

    for index, result in enumerate(results, start=1):
        result["rank"] = index

    if results:
        best_result = results[0]
        summary = {
            "total_resumes": len(results),
            "best_resume": best_result["file_name"],
            "highest_score": best_result["final_score"],
            "recommendation": best_result["recommendation"],
        }
    else:
        summary = {
            "total_resumes": 0,
            "best_resume": "",
            "highest_score": 0,
            "recommendation": "",
        }

    return {
        "results": results,
        "summary": summary,
    }
